#!/usr/bin/env python3
"""Generate the Home "New from template" starter files.

Builds a letter, a resume, a budget, an invoice and a short presentation,
and writes them base64-encoded into apps/shell/src/main/home-templates-data.ts
(lazily imported, so the bytes cost nothing until a template is picked).
Needs python-docx, openpyxl and python-pptx. Formula cells get their cached
values written in, so the workbooks show totals before any recalculation.

Usage: python3 tools/gen-home-templates.py [--files /path/out-dir]
"""
import base64
import io
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGB
from pptx.util import Inches as PptxInches, Pt as PptxPt

ACCENT = "1F4E79"
# Start from the app's own blank files (tools/gen-shell-new-templates.ts) so
# the templates carry its fonts, list numbering and 16:9 slide size.
BLANK = Path(__file__).resolve().parent.parent / "apps/shell/build/shell-new"


def deterministic(data: bytes) -> bytes:
    """Rewrite a package with fixed timestamps so regeneration is byte-stable."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                body = re.sub(rb"<dcterms:(created|modified)([^>]*)>[^<]*<",
                              rb"<dcterms:\1\2>2000-01-01T00:00:00Z<", body)
                body = re.sub(rb"<cp:lastModifiedBy>[^<]*<", rb"<cp:lastModifiedBy><", body)
                body = re.sub(rb"<dc:creator>[^<]*<", rb"<dc:creator><", body)
            info = zipfile.ZipInfo(item.filename, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, body)
    return out.getvalue()


def save_docx(doc, path: Path) -> None:
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(deterministic(buf.getvalue()))


def base_document():
    doc = Document(str(BLANK / "blank.docx"))
    body = doc.element.body
    for paragraph in body.findall(qn("w:p")):
        body.remove(paragraph)
    return doc


def bullet(doc, text: str) -> None:
    """A bullet on the blank file's numId 1 (abstractNum 0 is the bullet list)."""
    p = doc.add_paragraph(text, style="List Paragraph")
    num_pr = OxmlElement("w:numPr")
    for tag, value in (("w:ilvl", "0"), ("w:numId", "1")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), value)
        num_pr.append(el)
    p._p.get_or_add_pPr().append(num_pr)
    p.paragraph_format.space_after = Pt(0)


def letter(path: Path) -> None:
    doc = base_document()
    for line in ("[Your Name]", "[Street Address]", "[City, Postal Code]", "[Email] · [Phone]"):
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    doc.add_paragraph("[Date]")
    for line in ("[Recipient Name]", "[Title]", "[Company]", "[Street Address]", "[City, Postal Code]"):
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    doc.add_paragraph("Dear [Recipient Name],")
    doc.add_paragraph(
        "Start with a sentence that says why you are writing. Keep the first paragraph short "
        "so the reader knows straight away what the letter is about."
    )
    doc.add_paragraph(
        "Use the middle paragraphs for the details: what happened, what you need, and any "
        "dates or reference numbers the reader should have at hand."
    )
    doc.add_paragraph(
        "Close by saying what you would like to happen next and how the reader can reach you."
    )
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph()
    doc.add_paragraph("[Your Name]")
    save_docx(doc, path)


def resume(path: Path) -> None:
    doc = base_document()
    name = doc.add_paragraph()
    run = name.add_run("[Your Name]")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    name.paragraph_format.space_after = Pt(0)
    contact = doc.add_paragraph("[City] · [Email] · [Phone] · [Website or LinkedIn]")
    contact.paragraph_format.space_after = Pt(12)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor.from_string(ACCENT)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)

    def role(title: str, place: str, dates: str) -> None:
        p = doc.add_paragraph()
        p.add_run(title).bold = True
        p.add_run(f" · {place}")
        p.paragraph_format.space_after = Pt(0)
        d = doc.add_paragraph(dates)
        d.runs[0].italic = True
        d.paragraph_format.space_after = Pt(2)

    heading("Summary")
    doc.add_paragraph(
        "Two or three sentences on who you are, what you are good at, and the kind of role you want."
    )
    heading("Experience")
    role("[Job Title]", "[Organization, City]", "[Start] – [End]")
    for text in ("[An achievement, with a number if you have one]",
                 "[A responsibility that matters for the role you want]"):
        bullet(doc, text)
    role("[Job Title]", "[Organization, City]", "[Start] – [End]")
    for text in ("[An achievement]", "[A responsibility]"):
        bullet(doc, text)
    heading("Education")
    role("[Degree]", "[School, City]", "[Year]")
    heading("Skills")
    doc.add_paragraph("[Skill] · [Skill] · [Skill] · [Language]")
    save_docx(doc, path)


THIN = Side(style="thin", color="BFBFBF")
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '#,##0.00'


def save_xlsx(wb, path: Path, cached: dict) -> None:
    """Save, then write each formula cell's cached value next to its <f>."""
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = body.decode()
                for ref, value in cached.items():
                    pattern = re.compile(rf'(<c r="{ref}"[^>]*>)(<f>[^<]*</f>)(<v>[^<]*</v>|<v\s*/>)?')
                    xml, n = pattern.subn(lambda m: f"{m.group(1)}{m.group(2)}<v>{value:g}</v>", xml)
                    assert n == 1, ref
                body = xml.encode()
            dst.writestr(item, body)
    path.write_bytes(deterministic(out.getvalue()))


def style_header(ws, row: int, first: int, last: int) -> None:
    for col in range(first, last + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center" if col > first else "left")


def budget(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Monthly Budget"
    ws["A1"].font = Font(bold=True, size=18, color=ACCENT)
    ws.row_dimensions[1].height = 28
    ws["A2"] = "[Month, Year]"
    ws["A2"].font = Font(italic=True, color="595959")
    cached = {}

    def block(top: int, title: str, rows: list) -> int:
        ws.cell(row=top, column=1, value=title)
        ws.cell(row=top, column=2, value="Planned")
        ws.cell(row=top, column=3, value="Actual")
        ws.cell(row=top, column=4, value="Difference")
        style_header(ws, top, 1, 4)
        for i, (label, planned, actual) in enumerate(rows, start=top + 1):
            ws.cell(row=i, column=1, value=label)
            ws.cell(row=i, column=2, value=planned)
            ws.cell(row=i, column=3, value=actual)
            ws.cell(row=i, column=4, value=f"=C{i}-B{i}")
            cached[f"D{i}"] = actual - planned
        total = top + len(rows) + 1
        ws.cell(row=total, column=1, value="Total").font = Font(bold=True)
        sums = {}
        for col, letter in ((2, "B"), (3, "C"), (4, "D")):
            cell = ws.cell(row=total, column=col, value=f"=SUM({letter}{top + 1}:{letter}{total - 1})")
            cell.font = Font(bold=True)
            cell.border = Border(top=THIN)
        sums["B"] = sum(r[1] for r in rows)
        sums["C"] = sum(r[2] for r in rows)
        sums["D"] = sums["C"] - sums["B"]
        for letter, value in sums.items():
            cached[f"{letter}{total}"] = value
        for r in range(top + 1, total + 1):
            for c in (2, 3, 4):
                ws.cell(row=r, column=c).number_format = MONEY
        return total

    income_rows = [("Salary", 3000, 3000), ("Other income", 200, 150)]
    expense_rows = [("Rent", 1000, 1000), ("Utilities", 150, 165), ("Groceries", 400, 380),
                    ("Transport", 120, 140), ("Health", 80, 60), ("Savings", 500, 500),
                    ("Other", 200, 170)]
    income_total = block(4, "Income", income_rows)
    expense_top = income_total + 2
    expense_total = block(expense_top, "Expenses", expense_rows)
    left = expense_total + 2
    ws.cell(row=left, column=1, value="Left over").font = Font(bold=True, size=12)
    for col, letter in ((2, "B"), (3, "C")):
        cell = ws.cell(row=left, column=col, value=f"={letter}{income_total}-{letter}{expense_total}")
        cell.font = Font(bold=True, size=12)
        cell.number_format = MONEY
    cached[f"B{left}"] = sum(r[1] for r in income_rows) - sum(r[1] for r in expense_rows)
    cached[f"C{left}"] = sum(r[2] for r in income_rows) - sum(r[2] for r in expense_rows)
    ws.column_dimensions["A"].width = 24
    for letter in "BCD":
        ws.column_dimensions[letter].width = 14
    save_xlsx(wb, path, cached)


def invoice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = "INVOICE"
    ws["A1"].font = Font(bold=True, size=20, color=ACCENT)
    ws.row_dimensions[1].height = 30
    for row, (label, value) in enumerate((("Invoice no.", "[0001]"), ("Date", "[Date]"),
                                          ("Due", "[Due date]")), start=1):
        ws.cell(row=row, column=4, value=label).font = Font(bold=True)
        ws.cell(row=row, column=5, value=value)
    for row, text in enumerate(("[Your Business Name]", "[Address]", "[Email] · [Phone]"), start=3):
        ws.cell(row=row, column=1, value=text)
    ws["A7"] = "Bill to"
    ws["A7"].font = Font(bold=True, color=ACCENT)
    for row, text in enumerate(("[Client Name]", "[Client Address]"), start=8):
        ws.cell(row=row, column=1, value=text)
    head = 11
    for col, title in enumerate(("Description", "", "Qty", "Unit price", "Amount"), start=1):
        ws.cell(row=head, column=col, value=title or None)
    style_header(ws, head, 1, 5)
    items = [("[Item or service]", 2, 50), ("[Item or service]", 1, 120), ("[Item or service]", 3, 15)]
    cached = {}
    for i, (desc, qty, price) in enumerate(items, start=head + 1):
        ws.cell(row=i, column=1, value=desc)
        ws.cell(row=i, column=3, value=qty)
        ws.cell(row=i, column=4, value=price).number_format = MONEY
        cell = ws.cell(row=i, column=5, value=f"=C{i}*D{i}")
        cell.number_format = MONEY
        cached[f"E{i}"] = qty * price
    last = head + len(items)
    subtotal = sum(q * p for _, q, p in items)
    rate = 0.1
    rows = (("Subtotal", f"=SUM(E{head + 1}:E{last})", subtotal),
            ("Tax rate", None, rate),
            ("Tax", f"=E{last + 2}*E{last + 3}", subtotal * rate),
            ("Total", f"=E{last + 2}+E{last + 4}", subtotal * (1 + rate)))
    for offset, (label, formula, value) in enumerate(rows, start=2):
        r = last + offset
        ws.cell(row=r, column=4, value=label).font = Font(bold=True)
        cell = ws.cell(row=r, column=5, value=formula if formula else value)
        cell.number_format = "0%" if label == "Tax rate" else MONEY
        if formula:
            cached[f"E{r}"] = value
        if label == "Total":
            cell.font = Font(bold=True, size=12)
            cell.border = Border(top=THIN)
    ws.cell(row=last + 7, column=1, value="Thank you for your business.").font = Font(italic=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 4
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    save_xlsx(wb, path, cached)


def presentation(path: Path) -> None:
    prs = Presentation(str(BLANK / "blank.pptx"))
    accent = PptxRGB.from_string(ACCENT)
    blank_layout = prs.slide_layouts[0]
    width = prs.slide_width
    margin = PptxInches(0.9)

    def text(slide, top, height, lines, size, bold=False, color=None, gray=False):
        box = slide.shapes.add_textbox(margin, PptxInches(top), width - 2 * margin, PptxInches(height))
        frame = box.text_frame
        frame.word_wrap = True
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = line
            run.font.size = PptxPt(size)
            run.font.bold = bold
            if color is not None:
                run.font.color.rgb = color
            elif gray:
                run.font.color.rgb = PptxRGB.from_string("595959")
        return box

    def bar(slide, top):
        line = slide.shapes.add_shape(1, margin, PptxInches(top), PptxInches(1.2), PptxInches(0.08))
        line.fill.solid()
        line.fill.fore_color.rgb = accent
        line.line.fill.background()

    def slide():
        return prs.slides.add_slide(blank_layout)

    first = prs.slides[0] if len(prs.slides) else slide()
    text(first, 2.4, 1.3, ["[Presentation Title]"], 48, bold=True, color=accent)
    bar(first, 3.75)
    text(first, 4.0, 0.8, ["[Your Name] · [Date]"], 22, gray=True)

    agenda = slide()
    text(agenda, 0.6, 1.0, ["Agenda"], 36, bold=True, color=accent)
    text(agenda, 1.8, 4.5, ["1.  [First topic]", "2.  [Second topic]", "3.  [Third topic]",
                            "4.  [Questions]"], 24)

    section = slide()
    text(section, 2.8, 1.1, ["[Section Title]"], 40, bold=True, color=accent)
    bar(section, 3.95)
    text(section, 4.2, 0.8, ["[One line on what this section covers]"], 20, gray=True)

    content = slide()
    text(content, 0.6, 1.0, ["[Slide Title]"], 36, bold=True, color=accent)
    text(content, 1.8, 4.5, ["• [Key point]", "• [Supporting detail]", "• [Another point]"], 24)

    closing = slide()
    text(closing, 2.6, 1.1, ["Thank you"], 44, bold=True, color=accent)
    bar(closing, 3.75)
    text(closing, 4.0, 0.8, ["[Contact details or next steps]"], 22, gray=True)

    buf = io.BytesIO()
    prs.save(buf)
    path.write_bytes(deterministic(buf.getvalue()))


MODULE = BLANK.parent.parent / "src/main/home-templates-data.ts"
BUILDERS = (
    ("letter", "docx", letter),
    ("resume", "docx", resume),
    ("budget", "xlsx", budget),
    ("invoice", "xlsx", invoice),
    ("presentation", "pptx", presentation),
)


def main() -> None:
    files_dir = Path(sys.argv[2]) if len(sys.argv) > 2 and sys.argv[1] == "--files" else None
    lines = [
        "// Generated by tools/gen-home-templates.py; do not edit by hand.",
        "/** Home \"New from template\" starter files, base64-encoded */",
        "export const HOME_TEMPLATE_DATA = {",
    ]
    with __import__("tempfile").TemporaryDirectory() as scratch:
        for name, ext, build in BUILDERS:
            path = Path(scratch) / f"{name}.{ext}"
            build(path)
            data = path.read_bytes()
            if files_dir is not None:
                files_dir.mkdir(parents=True, exist_ok=True)
                (files_dir / path.name).write_bytes(data)
            lines.append(f"  {name}: '{base64.b64encode(data).decode()}',")
    lines.append("} as const")
    MODULE.write_text("\n".join(lines) + "\n")


if __name__ == "__main__":
    main()
